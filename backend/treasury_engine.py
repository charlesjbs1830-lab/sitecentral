"""
EDV Jr. - Motor de Tesouraria Avançada, CPQ Jurídico, DRE e Governança Financeira
Módulo corporativo para Empresa Júnior de Direito com base na Lei Federal nº 13.267/2016.
Centros de Custo por Caso Jurídico, Precificação Paramétrica de Honorários (CPQ),
Forecasting Econométrico com Sazonalidade Acadêmica, Travas de Mentor/OAB e Prestação de Contas.
"""

import sqlite3
import os
import io
import math
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional

DB_PATH = os.path.join(os.path.dirname(__file__), "auth.db")

def get_treasury_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

# ==============================================================================
# 1. CRIAÇÃO DE TABELAS & SEEDING IDEMPOTENTE
# ==============================================================================
def init_treasury_tables(conn=None):
    close_at_end = False
    if conn is None:
        conn = get_treasury_conn()
        close_at_end = True

    cursor = conn.cursor()

    # 1. Catálogo de Serviços Jurídicos Padronizados
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS cpq_servicos_juridicos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        area TEXT NOT NULL,
        nome TEXT NOT NULL,
        codigo TEXT UNIQUE NOT NULL,
        lead_time_semanas INTEGER NOT NULL,
        default_hh INTEGER NOT NULL,
        senioridade_recomendada TEXT NOT NULL,
        valor_hh_base REAL NOT NULL,
        complexidade_fator REAL DEFAULT 1.15,
        descricao TEXT NOT NULL,
        entregaveis_tecnicos TEXT NOT NULL,
        custas_estimadas REAL DEFAULT 0.0,
        margem_estatutaria_pct REAL DEFAULT 0.15
    );
    """)

    # 2. Histórico de Propostas Comerciais de Honorários (CPQ)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS cpq_propostas_juridicas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_nome TEXT NOT NULL,
        contato_nome TEXT,
        contato_email TEXT,
        servico_id INTEGER,
        servico_nome TEXT NOT NULL,
        area_juridica TEXT NOT NULL,
        senioridade TEXT NOT NULL,
        horas_pesquisa INTEGER NOT NULL,
        horas_redacao INTEGER NOT NULL,
        horas_revisao INTEGER NOT NULL,
        total_hh INTEGER NOT NULL,
        hourly_rate REAL NOT NULL,
        multiplicador_risco REAL NOT NULL,
        custas_inpi_cartorio REAL DEFAULT 0.0,
        subtotal_honorarios REAL NOT NULL,
        fundo_reinvestimento_estatutario REAL NOT NULL,
        preco_total REAL NOT NULL,
        status TEXT DEFAULT 'proposta_gerada',
        observacoes TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 3. Centros de Custo Jurídicos (DRE Gerencial por Projeto)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS dre_centros_custo_juridicos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        contrato_id INTEGER,
        caso_nome TEXT NOT NULL,
        cliente_nome TEXT NOT NULL,
        area_juridica TEXT NOT NULL,
        receita_bruta REAL NOT NULL,
        custas_diretas_inpi_cartorio REAL DEFAULT 0.0,
        despesas_operacionais_diretas REAL DEFAULT 0.0,
        margem_contribuicao REAL NOT NULL,
        margem_contribuicao_pct REAL NOT NULL,
        despesas_administrativas_rateadas REAL NOT NULL,
        superavit_estatutario REAL NOT NULL,
        reinvestimento_obrigatorio_pct REAL DEFAULT 1.0,
        fundo_capacitacao_reserva REAL NOT NULL,
        data_competencia TEXT DEFAULT '2026-03',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 4. Marcos Técnicos do Caso & Trava do Advogado/Professor Orientador
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS marcos_casos_juridicos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        centro_custo_id INTEGER NOT NULL,
        caso_nome TEXT NOT NULL,
        fase_nome TEXT NOT NULL,
        semana_inicio INTEGER NOT NULL,
        semana_fim INTEGER NOT NULL,
        hh_orcado INTEGER NOT NULL,
        hh_consumido REAL DEFAULT 0.0,
        progresso_pct REAL DEFAULT 0.0,
        is_caminho_critico INTEGER DEFAULT 0,
        requires_mentor_approval INTEGER DEFAULT 0,
        mentor_approval_status TEXT DEFAULT 'nao_requerido',
        mentor_nome TEXT,
        mentor_rubrica TEXT,
        mentor_parecer TEXT,
        data_aprovacao TEXT,
        status TEXT DEFAULT 'pendente',
        FOREIGN KEY (centro_custo_id) REFERENCES dre_centros_custo_juridicos(id) ON DELETE CASCADE
    );
    """)

    # 5. Pacotes do Orçamento Base Zero (OBZ)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS obz_pacotes_despesas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        pacote_nome TEXT NOT NULL,
        descricao TEXT NOT NULL,
        valor_planejado_anual REAL NOT NULL,
        valor_executado REAL NOT NULL,
        cor_hex TEXT NOT NULL
    );
    """)

    conn.commit()

    # Seeding do Catálogo de Serviços Jurídicos EDV
    cursor.execute("SELECT COUNT(*) FROM cpq_servicos_juridicos;")
    if cursor.fetchone()[0] == 0:
        servicos = [
            ("Propriedade Intelectual", "Registro de Marca no INPI", "EDV-PI-MARCA", 8, 45, "Consultor Jurídico", 55.00, 1.15,
             "Busca prévia de anterioridade, protocolo de pedido de registro perante o INPI, acompanhamento de RPI e defesa técnica.",
             "Relatório de Viabilidade de Marca, Parecer de Classificação de Nice, Protocolo Oficial INPI e Relatório de Auditoria RPI.", 142.00),
            ("Compliance & Regulação", "Adequação e Compliance LGPD", "EDV-CMP-LGPD", 10, 70, "Consultor Pleno", 75.00, 1.25,
             "Mapeamento de fluxo de dados pessoais, inventário de riscos, elaboração de RIPD, Termos de Uso e Política de Privacidade.",
             "Matriz de Riscos LGPD, Relatório de Impacto (RIPD), Política de Privacidade, Termo de Consentimento e Minuta de DPA com Fornecedores.", 0.0),
            ("Direito Contratual", "Elaboração e Revisão de Contratos Empresariais", "EDV-CTR-ELAB", 4, 35, "Consultor Jurídico", 55.00, 1.10,
             "Redação técnica customizada de instrumentos contratuais preventivos, cláusulas resolutivas, garantias e termos de confidencialidade (NDA).",
             "Minuta Contratual Customizada, Matriz de Riscos de Cláusulas, Guia de Boas Práticas de Execução e Versão Final Pronta para Assinatura.", 0.0),
            ("Direito Societário", "Acordo de Sócios & Planejamento Societário", "EDV-SOC-ACORD", 6, 50, "Consultor Pleno", 75.00, 1.20,
             "Estruturação de governança societária entre sócios/quotistas, regras de saída (tag along, drag along), vesting e resolução de conflitos.",
             "Minuta de Acordo de Sócios, Tabela de Vesting, Cláusulas de Não-Concorrência e Memorando de Entendimentos (MoU).", 0.0),
            ("Consultoria & Pareceres", "Parecer Jurídico & Viabilidade Regulatória", "EDV-CON-PARC", 5, 40, "Consultor Especialista", 85.00, 1.30,
             "Pesquisa aprofundada de doutrina e jurisprudência dos tribunais superiores (STJ/STF) para respaldar tomadas de decisão empresariais de alto risco.",
             "Parecer Jurídico Formal com Fundamentação Legal, Ementário Jurisprudencial, Análise Probabilística de Riscos e Recomendações Táticas.", 0.0),
            ("Terceiro Setor & EJs", "Estatuto Social e Registro em Cartório (RCPJ)", "EDV-ASS-ESTAT", 6, 45, "Consultor Jurídico", 55.00, 1.15,
             "Assessoria para estruturação estatutária de Associações sem fins lucrativos, Empresas Juniores e entidades esportivas/acadêmicas.",
             "Minuta de Estatuto Social Conforme Lei 13.267 / Código Civil, Ata de Fundação/Eleição, Edital de Convocação e Assessoria para Registro em Cartório.", 380.00)
        ]
        for s in servicos:
            cursor.execute("""
            INSERT INTO cpq_servicos_juridicos (
                area, nome, codigo, lead_time_semanas, default_hh, senioridade_recomendada,
                valor_hh_base, complexidade_fator, descricao, entregaveis_tecnicos, custas_estimadas, margem_estatutaria_pct
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0.15);
            """, s)

    # Seeding de Propostas CPQ
    cursor.execute("SELECT COUNT(*) FROM cpq_propostas_juridicas;")
    if cursor.fetchone()[0] == 0:
        propostas = [
            ("TechInovação Software Ltda.", "Lucas Silveira", "lucas@techinovacao.com.br", 2, "Adequação e Compliance LGPD", "Compliance & Regulação",
             "Consultor Pleno", 15, 40, 15, 70, 75.0, 1.25, 0.0, 6562.50, 984.38, 7546.88, "convertida", "Projeto prioritário para adequação da plataforma SaaS B2B."),
            ("Café Conilon Exportação S.A.", "Carlos Alberto Conilon", "carlos@cafeconilon.com.br", 1, "Registro de Marca no INPI", "Propriedade Intelectual",
             "Consultor Jurídico", 10, 25, 10, 45, 55.0, 1.15, 355.0, 2846.25, 426.94, 3628.19, "convertida", "Acompanhamento com prioridade de depósito e oposição de anterioridade."),
            ("FinTech Capixaba Ltda.", "Mariana Diniz", "mariana@capixabafin.com", 4, "Acordo de Sócios & Planejamento Societário", "Direito Societário",
             "Consultor Pleno", 10, 30, 10, 50, 75.0, 1.20, 0.0, 4500.0, 675.0, 5175.0, "em_negociacao", "Negociação de vesting e entrada de investidor anjo.")
        ]
        for p in propostas:
            cursor.execute("""
            INSERT INTO cpq_propostas_juridicas (
                cliente_nome, contato_nome, contato_email, servico_id, servico_nome, area_juridica,
                senioridade, horas_pesquisa, horas_redacao, horas_revisao, total_hh, hourly_rate,
                multiplicador_risco, custas_inpi_cartorio, subtotal_honorarios, fundo_reinvestimento_estatutario,
                preco_total, status, observacoes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, p)

    # Seeding de Centros de Custo (DRE) e seus Marcos com Trava de Mentor
    cursor.execute("SELECT COUNT(*) FROM dre_centros_custo_juridicos;")
    if cursor.fetchone()[0] == 0:
        casos = [
            ("Adequação LGPD Integral", "TechInovação Software Ltda.", "Compliance & Regulação",
             6200.00, 0.0, 310.00, 5890.00, 95.0, 620.00, 5270.00, 1.0, 5270.00, "2026-03"),
            ("Registro de Marca com Oposição INPI", "Café Conilon Exportação S.A.", "Propriedade Intelectual",
             3850.00, 355.00, 192.50, 3302.50, 85.8, 385.00, 2917.50, 1.0, 2917.50, "2026-03"),
            ("Acordo de Sócios e Vesting para Startup", "FinTech Capixaba Ltda.", "Direito Societário",
             4900.00, 0.0, 245.00, 4655.00, 95.0, 490.00, 4165.00, 1.0, 4165.00, "2026-02"),
            ("Revisão de Minutas de Prestação de Serviços", "Logística Vix Transportes", "Direito Contratual",
             2800.00, 0.0, 140.00, 2660.00, 95.0, 280.00, 2380.00, 1.0, 2380.00, "2026-02")
        ]
        for c in casos:
            cursor.execute("""
            INSERT INTO dre_centros_custo_juridicos (
                caso_nome, cliente_nome, area_juridica, receita_bruta, custas_diretas_inpi_cartorio,
                despesas_operacionais_diretas, margem_contribuicao, margem_contribuicao_pct,
                despesas_administrativas_rateadas, superavit_estatutario, reinvestimento_obrigatorio_pct,
                fundo_capacitacao_reserva, data_competencia
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, c)
            cc_id = cursor.lastrowid

            # Inserir marcos representativos com trava de mentor
            if cc_id == 1:
                marcos = [
                    (cc_id, c[0], "Fase 1: Mapeamento de Fluxos e Inventário de Dados", 1, 3, 20, 20.0, 100.0, 0, 0, "nao_requerido", None, None, None, None, "concluido"),
                    (cc_id, c[0], "Fase 2: Elaboração do RIPD e Políticas de Privacidade", 4, 7, 30, 26.0, 85.0, 1, 0, "nao_requerido", None, None, None, None, "em_andamento"),
                    (cc_id, c[0], "Fase 3: Validação Técnica do Advogado Orientador (Trava OAB)", 8, 9, 15, 0.0, 0.0, 1, 1, "pendente", None, None, None, None, "pendente"),
                    (cc_id, c[0], "Fase 4: Entrega Executiva e Treinamento de Encarregado (DPO)", 10, 10, 5, 0.0, 0.0, 0, 0, "nao_requerido", None, None, None, None, "pendente")
                ]
            elif cc_id == 2:
                marcos = [
                    (cc_id, c[0], "Fase 1: Busca de Anterioridade e Parecer de Risco", 1, 2, 15, 15.0, 100.0, 0, 0, "nao_requerido", None, None, None, None, "concluido"),
                    (cc_id, c[0], "Fase 2: Protocolo do Pedido e Pagamento GRU no INPI", 3, 4, 15, 15.0, 100.0, 1, 0, "nao_requerido", None, None, None, None, "concluido"),
                    (cc_id, c[0], "Fase 3: Análise de Oposição na RPI e Parecer do Mentor", 5, 8, 15, 8.0, 50.0, 1, 1, "pendente", None, None, None, None, "em_andamento")
                ]
            else:
                marcos = [
                    (cc_id, c[0], "Fase 1: Entrevista de Alinhamento e Diagnóstico da Minuta", 1, 2, 15, 15.0, 100.0, 0, 0, "nao_requerido", None, None, None, None, "concluido"),
                    (cc_id, c[0], "Fase 2: Redação da Minuta Preliminar e Matriz de Riscos", 3, 5, 25, 18.0, 72.0, 1, 0, "nao_requerido", None, None, None, None, "em_andamento"),
                    (cc_id, c[0], "Fase 3: Revisão e Chancela Formal do Advogado Orientador", 6, 6, 10, 0.0, 0.0, 1, 1, "pendente", None, None, None, None, "pendente")
                ]

            for m in marcos:
                cursor.execute("""
                INSERT INTO marcos_casos_juridicos (
                    centro_custo_id, caso_nome, fase_nome, semana_inicio, semana_fim,
                    hh_orcado, hh_consumido, progresso_pct, is_caminho_critico,
                    requires_mentor_approval, mentor_approval_status, mentor_nome, mentor_rubrica,
                    mentor_parecer, data_aprovacao, status
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
                """, m)

    # Seeding do Orçamento Base Zero (OBZ)
    cursor.execute("SELECT COUNT(*) FROM obz_pacotes_despesas;")
    if cursor.fetchone()[0] == 0:
        pacotes = [
            ("Capacitação, Cursos e Congressos MEJ (ENEJ/FEJESP)", "Formação acadêmica e profissional dos membros e assessores jurídicos.", 18000.00, 14500.00, "#059669"),
            ("Softwares, Pesquisa e Bibliotecas Jurídicas", "Assinaturas do Jusbrasil, Projuris, certificados digitais e tokens OAB.", 7200.00, 5800.00, "#3b82f6"),
            ("Despesas Administrativas, Sede e Encargos Institucionais", "Taxas bancárias, materiais de escritório, anuidade Brasil Júnior e cartório.", 6000.00, 4200.00, "#0b1d3a"),
            ("Fundo de Reserva de Emergência e Contingência", "Garantia de solvência e contingências operacionais (100% retido).", 12000.00, 12000.00, "#f59e0b")
        ]
        for p in pacotes:
            cursor.execute("""
            INSERT INTO obz_pacotes_despesas (
                pacote_nome, descricao, valor_planejado_anual, valor_executado, cor_hex
            ) VALUES (?, ?, ?, ?, ?);
            """, p)

    conn.commit()
    if close_at_end:
        conn.close()


# ==============================================================================
# 2. MOTOR DE CPQ JURÍDICO (CONFIGURE, PRICE, QUOTE)
# ==============================================================================
TAXAS_SENIORIDADE_JURIDICA = {
    "Trainee / Assessor": 35.00,
    "Consultor Jurídico": 55.00,
    "Consultor Pleno": 75.00,
    "Consultor Especialista": 85.00,
    "Diretor Jurídico": 120.00
}

def get_cpq_servicos_juridicos(area: Optional[str] = None) -> List[Dict[str, Any]]:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    if area and area != "ALL":
        cursor.execute("SELECT * FROM cpq_servicos_juridicos WHERE area = ? ORDER BY nome ASC;", (area,))
    else:
        cursor.execute("SELECT * FROM cpq_servicos_juridicos ORDER BY area, nome ASC;")
    rows = []
    for r in cursor.fetchall():
        d = dict(r)
        d["codigo_servico"] = d.get("codigo")
        d["macro_area"] = d.get("area")
        hh = d.get("default_hh", 40)
        d["horas_pesquisa_padrao"] = int(hh * 0.25)
        d["horas_redacao_padrao"] = int(hh * 0.50)
        d["horas_revisao_padrao"] = int(hh * 0.25)
        d["custas_inpi_padrao"] = d.get("custas_estimadas", 0.0)
        rows.append(d)
    conn.close()
    return rows

def calculate_cpq_juridico(
    servico_id: int,
    horas_pesquisa: int = 10,
    horas_redacao: int = 25,
    horas_revisao: int = 10,
    senioridade: str = "Consultor Jurídico",
    risco_fator: float = 1.15,
    custas_inpi_cartorio: float = 0.0,
    hourly_rate_override: Optional[float] = None
) -> Dict[str, Any]:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM cpq_servicos_juridicos WHERE id = ?;", (servico_id,))
    svc = cursor.fetchone()
    conn.close()

    if not svc:
        raise ValueError(f"Serviço jurídico ID {servico_id} não encontrado no catálogo.")

    hourly_rate = hourly_rate_override or TAXAS_SENIORIDADE_JURIDICA.get(senioridade, svc["valor_hh_base"])
    total_hh = int(horas_pesquisa + horas_redacao + horas_revisao)

    custo_base_honorarios = total_hh * hourly_rate
    subtotal_com_risco = custo_base_honorarios * float(risco_fator)
    fundo_reinvestimento = subtotal_com_risco * float(svc["margem_estatutaria_pct"] or 0.15)
    preco_total = subtotal_com_risco + fundo_reinvestimento + float(custas_inpi_cartorio)

    return {
        "servico_id": svc["id"],
        "servico_nome": svc["nome"],
        "codigo": svc["codigo"],
        "codigo_servico": svc["codigo"],
        "area_juridica": svc["area"],
        "macro_area": svc["area"],
        "senioridade": senioridade,
        "hourly_rate": round(hourly_rate, 2),
        "preco_hh_base": round(hourly_rate, 2),
        "horas_pesquisa": horas_pesquisa,
        "horas_redacao": horas_redacao,
        "horas_revisao": horas_revisao,
        "total_hh": total_hh,
        "horas_totais": total_hh,
        "multiplicador_risco": risco_fator,
        "fator_risco": risco_fator,
        "custas_inpi_cartorio": round(float(custas_inpi_cartorio), 2),
        "custas_diretas": round(float(custas_inpi_cartorio), 2),
        "custo_base_honorarios": round(custo_base_honorarios, 2),
        "subtotal_hh": round(custo_base_honorarios, 2),
        "subtotal_honorarios": round(subtotal_com_risco, 2),
        "honorarios_brutos": round(subtotal_com_risco, 2),
        "fundo_reinvestimento_estatutario": round(fundo_reinvestimento, 2),
        "margem_estatutaria_reinvestimento": round(fundo_reinvestimento, 2),
        "preco_total_proposta": round(preco_total, 2),
        "conformidade_lei_13267": True,
        "margem_estatutaria_pct": round(svc["margem_estatutaria_pct"] * 100, 1)
    }

def save_proposta_cpq_juridica(data: Dict[str, Any]) -> Dict[str, Any]:
    conn = get_treasury_conn()
    cursor = conn.cursor()

    cliente_nome = data.get("cliente_nome", "").strip()
    if not cliente_nome:
        raise ValueError("O nome do cliente é obrigatório para registrar a proposta.")

    calc = calculate_cpq_juridico(
        servico_id=int(data.get("servico_id", 1)),
        horas_pesquisa=int(data.get("horas_pesquisa", 10)),
        horas_redacao=int(data.get("horas_redacao", 25)),
        horas_revisao=int(data.get("horas_revisao", 10)),
        senioridade=data.get("senioridade", "Consultor Jurídico"),
        risco_fator=float(data.get("multiplicador_risco", 1.15)),
        custas_inpi_cartorio=float(data.get("custas_inpi_cartorio", 0.0)),
        hourly_rate_override=data.get("hourly_rate")
    )

    cursor.execute("""
    INSERT INTO cpq_propostas_juridicas (
        cliente_nome, contato_nome, contato_email, servico_id, servico_nome,
        area_juridica, senioridade, horas_pesquisa, horas_redacao, horas_revisao,
        total_hh, hourly_rate, multiplicador_risco, custas_inpi_cartorio,
        subtotal_honorarios, fundo_reinvestimento_estatutario, preco_total,
        status, observacoes
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """, (
        cliente_nome,
        data.get("contato_nome", ""),
        data.get("contato_email", ""),
        calc["servico_id"],
        calc["servico_nome"],
        calc["area_juridica"],
        calc["senioridade"],
        calc["horas_pesquisa"],
        calc["horas_redacao"],
        calc["horas_revisao"],
        calc["total_hh"],
        calc["hourly_rate"],
        calc["multiplicador_risco"],
        calc["custas_inpi_cartorio"],
        calc["subtotal_honorarios"],
        calc["fundo_reinvestimento_estatutario"],
        calc["preco_total_proposta"],
        data.get("status", "proposta_gerada"),
        data.get("observacoes", "")
    ))
    new_id = cursor.lastrowid
    conn.commit()
    conn.close()

    return {"status": "success", "id": new_id, "proposta_id": new_id, "proposta": calc}

def list_propostas_cpq_juridicas() -> List[Dict[str, Any]]:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM cpq_propostas_juridicas ORDER BY id DESC;")
    rows = []
    for r in cursor.fetchall():
        d = dict(r)
        d["valor_total"] = d.get("preco_total", 0.0)
        d["margem_reinvestimento"] = d.get("fundo_reinvestimento_estatutario", 0.0)
        d["horas_totais"] = d.get("total_hh", 0)
        d["data_proposta"] = d.get("created_at")
        rows.append(d)
    conn.close()
    return rows

def update_proposta_cpq_status(proposta_id: int, status: str) -> bool:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("""
    UPDATE cpq_propostas_juridicas 
    SET status = ?, updated_at = CURRENT_TIMESTAMP
    WHERE id = ?;
    """, (status, proposta_id))
    affected = cursor.rowcount
    conn.commit()
    conn.close()
    return affected > 0

def delete_proposta_cpq(proposta_id: int) -> bool:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM cpq_propostas_juridicas WHERE id = ?;", (proposta_id,))
    affected = cursor.rowcount
    conn.commit()
    conn.close()
    return affected > 0


# ==============================================================================
# 3. DRE POR CENTRO DE CUSTO JURÍDICO & MARGENS
# ==============================================================================
def get_dre_juridica_centros_custo() -> Dict[str, Any]:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM dre_centros_custo_juridicos ORDER BY id ASC;")
    rows = []
    for r in cursor.fetchall():
        d = dict(r)
        d["margem_contribuicao_direta"] = d.get("margem_contribuicao", 0.0)
        d["superavit_estatutario_retido"] = d.get("superavit_estatutario", 0.0)
        rows.append(d)

    total_receita = sum(r["receita_bruta"] for r in rows)
    total_custas = sum(r["custas_diretas_inpi_cartorio"] + r["despesas_operacionais_diretas"] for r in rows)
    total_margem = sum(r["margem_contribuicao"] for r in rows)
    total_rateio = sum(r["despesas_administrativas_rateadas"] for r in rows)
    total_superavit = sum(r["superavit_estatutario"] for r in rows)
    margem_media_pct = round((total_margem / total_receita * 100), 1) if total_receita else 0.0

    # Anexar marcos de cada caso
    for r in rows:
        cursor.execute("SELECT * FROM marcos_casos_juridicos WHERE centro_custo_id = ? ORDER BY semana_inicio ASC;", (r["id"],))
        marcos_list = []
        for m in cursor.fetchall():
            md = dict(m)
            md["validado_orientador"] = 1 if (md.get("mentor_approval_status") == "aprovado" or md.get("data_aprovacao")) else 0
            marcos_list.append(md)
        r["marcos"] = marcos_list

    conn.close()

    consolidado = {
        "total_receita_bruta": round(total_receita, 2),
        "total_custas_diretas": round(total_custas, 2),
        "total_margem_contribuicao": round(total_margem, 2),
        "margem_contribuicao_media_pct": margem_media_pct,
        "total_despesas_rateadas": round(total_rateio, 2),
        "total_superavit_estatutario": round(total_superavit, 2),
        "reinvestimento_obrigatorio_pct": 100.0,
        "fundo_capacitacao_acumulado": round(total_superavit, 2),
        "vedacao_distribuicao_lucro": True,
        "lei_13267_compliance": "100% dos superávits destinados ao fundo estatutário de reserva e capacitação; R$ 0,00 de remuneração a membros."
    }

    return {
        "centros_custo": rows,
        "consolidado": consolidado,
        "receita_bruta_total": round(total_receita, 2),
        "custas_diretas_total": round(total_custas, 2),
        "margem_contribuicao_total": round(total_margem, 2),
        "margem_contribuicao_media_pct": margem_media_pct,
        "despesas_rateadas_total": round(total_rateio, 2),
        "superavit_estatutario_total": round(total_superavit, 2),
    }

def create_dre_centro_custo_juridico(data: Dict[str, Any]) -> Dict[str, Any]:
    conn = get_treasury_conn()
    cursor = conn.cursor()

    caso_nome = data.get("caso_nome", "").strip()
    cliente_nome = data.get("cliente_nome", "").strip()
    area = data.get("area_juridica", "Direito Contratual")
    receita_bruta = float(data.get("receita_bruta", 4500.0))
    custas_inpi = float(data.get("custas_diretas_inpi_cartorio") or 0.0)
    desp_param = data.get("despesas_operacionais_diretas")
    desp_diretas = float(desp_param) if desp_param is not None else round(receita_bruta * 0.03, 2)

    margem = receita_bruta - custas_inpi - desp_diretas
    margem_pct = round((margem / receita_bruta * 100), 1) if receita_bruta else 0.0
    rateio_adm = round(receita_bruta * 0.10, 2)
    superavit = margem - rateio_adm

    cursor.execute("""
    INSERT INTO dre_centros_custo_juridicos (
        caso_nome, cliente_nome, area_juridica, receita_bruta, custas_diretas_inpi_cartorio,
        despesas_operacionais_diretas, margem_contribuicao, margem_contribuicao_pct,
        despesas_administrativas_rateadas, superavit_estatutario, reinvestimento_obrigatorio_pct,
        fundo_capacitacao_reserva, data_competencia
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1.0, ?, ?);
    """, (
        caso_nome, cliente_nome, area, receita_bruta, custas_inpi,
        desp_diretas, margem, margem_pct, rateio_adm, superavit, superavit,
        data.get("data_competencia", datetime.now().strftime("%Y-%m"))
    ))
    new_id = cursor.lastrowid

    # Criar 3 marcos padrão para o caso com trava de orientador
    marcos_padrao = [
        (new_id, caso_nome, "Alinhamento de Escopo & Coleta de Documentos", 1, 2, 15, 0.0, 0.0, 0, 0, "nao_requerido", "em_andamento"),
        (new_id, caso_nome, "Pesquisa Jurisprudencial e Redação da Minuta Preliminar", 3, 5, 25, 0.0, 0.0, 1, 0, "nao_requerido", "pendente"),
        (new_id, caso_nome, "Validação e Parecer Técnico do Advogado Orientador (Trava OAB)", 6, 6, 10, 0.0, 0.0, 1, 1, "pendente", "pendente")
    ]
    for m in marcos_padrao:
        cursor.execute("""
        INSERT INTO marcos_casos_juridicos (
            centro_custo_id, caso_nome, fase_nome, semana_inicio, semana_fim,
            hh_orcado, hh_consumido, progresso_pct, is_caminho_critico,
            requires_mentor_approval, mentor_approval_status, status
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, m)

    conn.commit()
    conn.close()

    return {"status": "success", "id": new_id, "caso_nome": caso_nome}

def aprovar_marco_orientador_juridico(marco_id: int, mentor_nome: str, mentor_rubrica: str, mentor_parecer: str) -> bool:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("""
    UPDATE marcos_casos_juridicos
    SET mentor_approval_status = 'aprovado',
        mentor_nome = ?,
        mentor_rubrica = ?,
        mentor_parecer = ?,
        data_aprovacao = CURRENT_TIMESTAMP,
        progresso_pct = 100.0,
        status = 'concluido'
    WHERE id = ?;
    """, (mentor_nome, mentor_rubrica, mentor_parecer, marco_id))
    affected = cursor.rowcount
    conn.commit()
    conn.close()
    return affected > 0


# ==============================================================================
# 4. MOTOR ECONOMÉTRICO DE FORECASTING & RUNWAY JURÍDICO (12 A 36 MESES)
# ==============================================================================
def get_forecasting_juridico(
    horizon_months: int = 24,
    growth_rate_pct: float = 15.0,
    opt_bonus_pct: float = 20.0,
    pess_penalty_pct: float = 15.0,
    inflation_pct: float = 5.5,
    seasonality_intensity: float = 1.0,
    initial_cash_balance: Optional[float] = None
) -> Dict[str, Any]:
    conn = get_treasury_conn()
    cursor = conn.cursor()

    if initial_cash_balance is None:
        cursor.execute("""
        SELECT COALESCE(SUM(CASE WHEN tipo = 'receita' AND status = 'pago' THEN valor ELSE 0 END), 0.0) -
               COALESCE(SUM(CASE WHEN tipo = 'despesa' AND status = 'pago' THEN valor ELSE 0 END), 0.0) as saldo
        FROM transacoes_financeiras;
        """)
        row = cursor.fetchone()
        calc_saldo = float(row["saldo"]) if row and row["saldo"] is not None else 0.0
        initial_cash_balance = calc_saldo if calc_saldo > 1000 else 38500.0

    conn.close()

    months = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
    now = datetime.now()
    labels = []
    cash_base = []
    cash_opt = []
    cash_pess = []
    rev_base = []
    rev_opt = []
    rev_pess = []

    cur_base = initial_cash_balance
    cur_opt = initial_cash_balance
    cur_pess = initial_cash_balance

    base_monthly_revenue = 9800.0
    base_monthly_cost = 3200.0

    # Sazonalidade MEJ Direito (Férias/OAB em Jan/Jul/Dez, Alta em Mar/Abr/Mai/Set/Out)
    seasonal_weights = [0.60, 0.80, 1.25, 1.30, 1.20, 1.05, 0.70, 0.95, 1.25, 1.20, 1.00, 0.70]

    for m in range(1, horizon_months + 1):
        target_month_idx = (now.month - 1 + m) % 12
        year_offset = (now.month - 1 + m) // 12
        year = now.year + year_offset
        month_label = f"{months[target_month_idx]}/{str(year)[-2:]}"
        labels.append(month_label)

        s_factor = 1.0 + (seasonal_weights[target_month_idx] - 1.0) * seasonality_intensity

        monthly_growth = math.pow(1.0 + growth_rate_pct / 100.0, 1 / 12) - 1.0
        monthly_inflation = math.pow(1.0 + inflation_pct / 100.0, 1 / 12) - 1.0

        # Base
        r_b = base_monthly_revenue * math.pow(1.0 + monthly_growth, m) * s_factor
        c_b = base_monthly_cost * math.pow(1.0 + monthly_inflation, m)
        cur_base += (r_b - c_b)
        rev_base.append(round(r_b))
        cash_base.append(round(cur_base))

        # Otimista (+20% captação contratos)
        r_o = r_b * (1.0 + opt_bonus_pct / 100.0)
        c_o = c_b * 0.95
        cur_opt += (r_o - c_o)
        rev_opt.append(round(r_o))
        cash_opt.append(round(cur_opt))

        # Pessimista (-15% captação + custos extras)
        r_p = r_b * (1.0 - pess_penalty_pct / 100.0)
        c_p = c_b * 1.10
        cur_pess += (r_p - c_p)
        rev_pess.append(round(r_p))
        cash_pess.append(round(cur_pess))

    # Cálculo de Runway (meses de cobertura de custos sem novas receitas)
    monthly_burn_rate = base_monthly_cost * 1.05
    runway_months = round(initial_cash_balance / monthly_burn_rate, 1) if monthly_burn_rate else 99.0

    projecoes = []
    for i, mes in enumerate(labels):
        projecoes.append({
            "mes": mes,
            "caixa_acumulado_base": cash_base[i],
            "caixa_acumulado_otimista": cash_opt[i],
            "caixa_acumulado_pessimista": cash_pess[i],
            "receita_base": rev_base[i],
            "receita_otimista": rev_opt[i],
            "receita_pessimista": rev_pess[i]
        })

    resumo_executivo = {
        "caixa_final_cenario_base": cash_base[-1],
        "caixa_final_cenario_otimista": cash_opt[-1],
        "caixa_final_cenario_pessimista": cash_pess[-1],
        "cash_runway_estimado_meses": runway_months,
        "receita_acumulada_base": sum(rev_base),
        "reinvestimento_estatutario_acumulado": round(sum(rev_base) * 0.85)
    }

    return {
        "metadata": {
            "horizon_months": horizon_months,
            "initial_cash_balance": initial_cash_balance,
            "growth_rate_annual_pct": growth_rate_pct,
            "inflation_annual_pct": inflation_pct,
            "runway_atual_meses": runway_months
        },
        "projecoes": projecoes,
        "resumo_executivo": resumo_executivo,
        "charts": {
            "time_labels": labels,
            "cash_balance": {
                "base": cash_base,
                "optimistic": cash_opt,
                "pessimistic": cash_pess
            },
            "monthly_revenue": {
                "base": rev_base,
                "optimistic": rev_opt,
                "pessimistic": rev_pess
            }
        },
        "summary": resumo_executivo
    }


# ==============================================================================
# 5. ORÇAMENTO BASE ZERO (OBZ) & PRESTAÇÃO DE CONTAS
# ==============================================================================
def get_obz_distribuicao_edv() -> Dict[str, Any]:
    conn = get_treasury_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM obz_pacotes_despesas ORDER BY id ASC;")
    pacotes = []
    for p in cursor.fetchall():
        d = dict(p)
        d["nome_pacote"] = d.get("pacote_nome")
        d["valor_orcado"] = d.get("valor_planejado_anual")
        d["valor_realizado"] = d.get("valor_executado")
        d["saldo_remanescente"] = round(d["valor_orcado"] - d["valor_realizado"], 2)
        d["macro_area"] = d.get("descricao", "Gestão Estratégica")
        pacotes.append(d)
    conn.close()

    total_planejado = sum(p["valor_orcado"] for p in pacotes)
    total_executado = sum(p["valor_realizado"] for p in pacotes)
    saldo_remanescente = round(total_planejado - total_executado, 2)
    exec_pct = round((total_executado / total_planejado * 100), 1) if total_planejado else 0.0

    return {
        "pacotes": pacotes,
        "total_orcado": round(total_planejado, 2),
        "total_planejado": round(total_planejado, 2),
        "total_realizado": round(total_executado, 2),
        "total_executado": round(total_executado, 2),
        "saldo_remanescente": saldo_remanescente,
        "execucao_global_pct": exec_pct,
        "taxa_execucao_pct": exec_pct,
        "labels": [p["pacote_nome"] for p in pacotes],
        "valores": [p["valor_executado"] for p in pacotes],
        "cores": [p["cor_hex"] for p in pacotes]
    }


# ==============================================================================
# 6. GERADOR DE RELATÓRIO DE PRESTAÇÃO DE CONTAS EXECUTIVO (EXCEL & HTML)
# ==============================================================================
def generate_prestacao_contas_excel_edv() -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    dre = get_dre_juridica_centros_custo()
    obz = get_obz_distribuicao_edv()

    wb = openpyxl.Workbook()

    # Aba 1: DRE por Centro de Custo
    ws1 = wb.active
    ws1.title = "DRE por Centro de Custo"
    ws1.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0B1D3A", end_color="0B1D3A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    accent_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)

    ws1.merge_cells("A1:H1")
    ws1["A1"] = "EDV JR. | RELATÓRIO DE DRE GERENCIAL POR CENTRO DE CUSTO (LEI Nº 13.267/2016)"
    ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws1["A1"].fill = header_fill
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 35

    cols = ["Caso / Projeto Jurídico", "Cliente", "Área de Foco", "Receita Bruta (R$)", "Custas Diretas (R$)", "Margem Contrib. (R$)", "Margem (%)", "Fundo Estatutário (R$)"]
    for col_idx, h in enumerate(cols, 1):
        cell = ws1.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[3].height = 25

    row_idx = 4
    for cc in dre["centros_custo"]:
        ws1.cell(row=row_idx, column=1, value=cc["caso_nome"])
        ws1.cell(row=row_idx, column=2, value=cc["cliente_nome"])
        ws1.cell(row=row_idx, column=3, value=cc["area_juridica"])
        ws1.cell(row=row_idx, column=4, value=cc["receita_bruta"]).number_format = 'R$ #,##0.00'
        ws1.cell(row=row_idx, column=5, value=cc["custas_diretas_inpi_cartorio"]).number_format = 'R$ #,##0.00'
        ws1.cell(row=row_idx, column=6, value=cc["margem_contribuicao"]).number_format = 'R$ #,##0.00'
        ws1.cell(row=row_idx, column=7, value=f"{cc['margem_contribuicao_pct']}%").alignment = Alignment(horizontal="center")
        ws1.cell(row=row_idx, column=8, value=cc["fundo_capacitacao_reserva"]).number_format = 'R$ #,##0.00'
        row_idx += 1

    # Linha Consolidada
    ws1.cell(row=row_idx, column=1, value="CONSOLIDADO EDV").font = bold_font
    ws1.cell(row=row_idx, column=4, value=dre["consolidado"]["total_receita_bruta"]).font = bold_font
    ws1.cell(row=row_idx, column=4).number_format = 'R$ #,##0.00'
    ws1.cell(row=row_idx, column=5, value=dre["consolidado"]["total_custas_diretas"]).font = bold_font
    ws1.cell(row=row_idx, column=5).number_format = 'R$ #,##0.00'
    ws1.cell(row=row_idx, column=6, value=dre["consolidado"]["total_margem_contribuicao"]).font = bold_font
    ws1.cell(row=row_idx, column=6).number_format = 'R$ #,##0.00'
    ws1.cell(row=row_idx, column=7, value=f"{dre['consolidado']['margem_contribuicao_media_pct']}%").font = bold_font
    ws1.cell(row=row_idx, column=7).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=8, value=dre["consolidado"]["total_superavit_estatutario"]).font = bold_font
    ws1.cell(row=row_idx, column=8).number_format = 'R$ #,##0.00'

    # Ajuste de largura de colunas
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Aba 2: Orçamento Base Zero (OBZ)
    ws2 = wb.create_sheet(title="Orçamento Base Zero")
    ws2.views.sheetView[0].showGridLines = True
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "DISTRIBUIÇÃO POR PACOTES OBZ & PRESTAÇÃO DE CONTAS AGO"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws2["A1"].fill = header_fill
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")

    cols2 = ["Pacote Orçamentário", "Descrição e Finalidade", "Planejado Anual (R$)", "Executado (R$)"]
    for col_idx, h in enumerate(cols2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = PatternFill(start_color="0B1D3A", end_color="0B1D3A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    r2 = 4
    for p in obz["pacotes"]:
        ws2.cell(row=r2, column=1, value=p["pacote_nome"])
        ws2.cell(row=r2, column=2, value=p["descricao"])
        ws2.cell(row=r2, column=3, value=p["valor_planejado_anual"]).number_format = 'R$ #,##0.00'
        ws2.cell(row=r2, column=4, value=p["valor_executado"]).number_format = 'R$ #,##0.00'
        r2 += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 16)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def generate_prestacao_contas_html_edv() -> str:
    dre = get_dre_juridica_centros_custo()
    obz = get_obz_distribuicao_edv()

    rows_html = "".join([
        f"""
        <tr>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; font-weight: bold; color: #0f172a;">{cc['caso_nome']}</td>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; color: #475569;">{cc['cliente_nome']}</td>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; font-size: 11px;">{cc['area_juridica']}</td>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; text-align: right; font-family: monospace;">R$ {cc['receita_bruta']:,.2f}</td>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; text-align: right; font-family: monospace; color: #e11d48;">R$ {cc['custas_diretas_inpi_cartorio']:,.2f}</td>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; text-align: right; font-family: monospace; font-weight: bold; color: #059669;">R$ {cc['margem_contribuicao']:,.2f} ({cc['margem_contribuicao_pct']}%)</td>
            <td style="padding: 10px; border-bottom: 1px solid #e2e8f0; text-align: right; font-family: monospace; font-weight: bold; color: #0b1d3a;">R$ {cc['fundo_capacitacao_reserva']:,.2f}</td>
        </tr>
        """ for cc in dre["centros_custo"]
    ])

    obz_html = "".join([
        f"""
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 12px; margin-bottom: 8px;">
            <div style="display: flex; justify-content: space-between; font-weight: bold; color: #0f172a;">
                <span>{p['pacote_nome']}</span>
                <span style="font-family: monospace; color: #059669;">R$ {p['valor_executado']:,.2f}</span>
            </div>
            <p style="margin: 4px 0 0 0; font-size: 11px; color: #64748b;">{p['descricao']}</p>
        </div>
        """ for p in obz["pacotes"]
    ])

    return f"""
    <!DOCTYPE html>
    <html lang="pt-BR">
    <head>
        <meta charset="UTF-8">
        <title>Prestação de Contas Executiva - EDV Jr.</title>
        <style>
            body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; margin: 30px; color: #334155; line-height: 1.5; }}
            .header {{ border-bottom: 3px solid #059669; padding-bottom: 15px; margin-bottom: 25px; }}
            .header h1 {{ margin: 0; color: #0b1d3a; font-size: 22px; }}
            .header p {{ margin: 5px 0 0 0; color: #64748b; font-size: 12px; }}
            .card {{ background: #ffffff; border: 1px solid #cbd5e1; border-radius: 10px; padding: 20px; margin-bottom: 25px; }}
            table {{ width: 100%; border-collapse: collapse; font-size: 12px; }}
            th {{ background: #0b1d3a; color: #ffffff; padding: 10px; text-align: left; }}
            .legal-box {{ background: #ecfdf5; border-left: 4px solid #059669; padding: 12px; border-radius: 4px; font-size: 11px; color: #065f46; margin-top: 20px; }}
        </style>
    </head>
    <body>
        <div class="header">
            <h1>EDV JR. | CONSULTORIA JURÍDICA</h1>
            <p>Relatório Oficial de Prestação de Contas, DRE e Governança Estatutária (Lei nº 13.267/2016)</p>
        </div>

        <div class="card">
            <h2 style="font-size: 15px; color: #0b1d3a; margin-top: 0;">1. DRE Gerencial por Centro de Custo Jurídico</h2>
            <table>
                <thead>
                    <tr>
                        <th>Caso Jurídico</th>
                        <th>Cliente</th>
                        <th>Área</th>
                        <th style="text-align: right;">Receita</th>
                        <th style="text-align: right;">Custas</th>
                        <th style="text-align: right;">Margem Direta</th>
                        <th style="text-align: right;">Superávit 100% Reinvestido</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                </tbody>
            </table>
            <div class="legal-box">
                <strong>Cláusula de Blindagem Estatutária:</strong> Em estrita conformidade com o Art. 2º e Art. 4º da Lei Federal nº 13.267/2016, 100% dos superávits operacionais apurados (R$ {dre['consolidado']['total_superavit_estatutario']:,.2f}) foram creditados no Fundo de Reserva e Capacitação Docente/Discente, sendo vedada qualquer distribuição a membros.
            </div>
        </div>

        <div class="card">
            <h2 style="font-size: 15px; color: #0b1d3a; margin-top: 0;">2. Distribuição por Pacotes do Orçamento Base Zero (OBZ)</h2>
            {obz_html}
        </div>
    </body>
    </html>
    """
